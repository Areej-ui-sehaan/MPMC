#!/usr/bin/env python3
"""
Compact Markdown -> PDF + DOCX converter for the MPMC Module-2 notes.

Keeps EVERYTHING from the markdown but squeezes the page count by using
  - A4 with narrow margins, small fonts, tight leading, minimal spacing
  - table typography whose size adapts to the number of columns
  - compact monospace code blocks
  - headings with keepWithNext (no page breaks, no orphan headings)

Usage:  python3 tools/notes_to_pdf_docx.py FILE.md [FILE2.md ...]
Output: <basename>.pdf and <basename>.docx next to each source file.
"""

import os
import re
import sys

# --------------------------------------------------------------------------- #
# Markdown parsing
# --------------------------------------------------------------------------- #

RE_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
RE_HRULE = re.compile(r"^-{3,}\s*$")
RE_TABLE_SEP = re.compile(r"^\|[\s:\-|]+\|$")
RE_UL = re.compile(r"^(\s*)[-*+]\s+(.*)$")
RE_OL = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
RE_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")
RE_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def _is_block_start(line: str) -> bool:
    s = line.strip()
    if not s:
        return True
    if RE_HEADING.match(line) or s.startswith("```") or s.startswith("|") or s.startswith(">"):
        return True
    if RE_UL.match(line) or RE_OL.match(line) or RE_HRULE.match(s):
        return True
    return False


def parse_blocks(text: str):
    """Markdown -> list of (kind, payload) blocks."""
    lines = text.replace("\r\n", "\n").split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if RE_HRULE.match(s):
            i += 1
            continue
        m = RE_HEADING.match(line)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("code", code))
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                i += 1
                if RE_TABLE_SEP.match(raw):
                    continue
                cells = [c.strip() for c in raw.strip("|").replace("\\|", "\x00").split("|")]
                rows.append([c.replace("\x00", "|") for c in cells])
            if rows:
                blocks.append(("table", rows))
            continue
        if s.startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("quote", quote))
            continue
        m = RE_UL.match(line)
        if m:
            items = []
            while i < n:
                mm = RE_UL.match(lines[i])
                if mm:
                    items.append((len(mm.group(1)), mm.group(2).strip()))
                    i += 1
                elif not lines[i].strip() and i + 1 < n and RE_UL.match(lines[i + 1]):
                    i += 1
                else:
                    break
            blocks.append(("ul", items))
            continue
        m = RE_OL.match(line)
        if m:
            items = []
            while i < n:
                mm = RE_OL.match(lines[i])
                if mm:
                    items.append((len(mm.group(1)), mm.group(3).strip()))
                    i += 1
                elif not lines[i].strip() and i + 1 < n and RE_OL.match(lines[i + 1]):
                    i += 1
                else:
                    break
            blocks.append(("ol", items))
            continue
        para = [s]
        i += 1
        while i < n and not _is_block_start(lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append(("p", " ".join(para)))
    return blocks


def inline_parts(text: str):
    """-> [(text, bold, code, italic), ...]"""
    text = RE_LINK.sub(r"\1", text)
    out, pos = [], 0
    for m in RE_INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False, False))
        tok = m.group(0)
        if tok.startswith("**"):
            out.append((tok[2:-2], True, False, False))
        elif tok.startswith("`"):
            out.append((tok[1:-1], False, True, False))
        else:
            out.append((tok[1:-1], False, False, True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False, False))
    return [(t, b, c, it) for (t, b, c, it) in out if t]


def plain(text: str) -> str:
    return "".join(t for t, *_ in inline_parts(text))



def _col_widths(rows, ncols, fs_pt, avail_pt, pad_pt=5.5):
    """Column widths (points) that fit `avail_pt`.

    Natural width = longest cell content; minimum width = longest single word,
    so a returned layout never breaks a word.  Returns None if the table cannot
    fit at this font size even at minimum widths (caller then reduces the font).
    """
    cw_nat = 0.55 * fs_pt + 0.75      # average advance of mixed-case text (pt/char)
    cw_min = 0.74 * fs_pt + 0.75      # worst case: all-caps / digits / symbols
    nat, mins = [], []
    for ci in range(ncols):
        longest_cell = max(len(plain(r[ci])) for r in rows) or 1
        longest_word = max((len(w) for r in rows for w in plain(r[ci]).split()), default=1)
        nat.append(longest_cell * cw_nat + pad_pt)
        mins.append(min(longest_word * cw_min + pad_pt, avail_pt))
    if sum(mins) > avail_pt:
        return None
    w = nat[:]
    total = sum(w)
    if total < avail_pt:                       # stretch unused space proportionally
        k = avail_pt / total
        return [x * k for x in w]
    for _ in range(300):                       # shrink, never below a minimum
        total = sum(w)
        if total <= avail_pt + 0.05:
            break
        slack = [max(0.0, w[i] - mins[i]) for i in range(ncols)]
        slack_total = sum(slack)
        if slack_total <= 0.01:
            return None
        take = min(total - avail_pt, slack_total)
        w = [w[i] - take * slack[i] / slack_total for i in range(ncols)]
    return w


# --------------------------------------------------------------------------- #
# PDF (reportlab)
# --------------------------------------------------------------------------- #

FONT_DIR = "/usr/share/fonts/truetype/dejavu/"


def _register_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    files = {
        "DJV": "DejaVuSans.ttf",
        "DJV-B": "DejaVuSans-Bold.ttf",
        "DJV-M": "DejaVuSansMono.ttf",
        "DJV-MB": "DejaVuSansMono-Bold.ttf",
    }
    for name, fn in files.items():
        pdfmetrics.registerFont(TTFont(name, FONT_DIR + fn))
    # family with real bold, italics fall back to the regular face (no oblique shipped)
    pdfmetrics.registerFontFamily("DJV", normal="DJV", bold="DJV-B", italic="DJV", boldItalic="DJV-B")
    # self-families so that every custom font name resolves (ps2tt lookups)
    for name in ("DJV-B", "DJV-M", "DJV-MB"):
        pdfmetrics.registerFontFamily(name, normal=name, bold=name, italic=name, boldItalic=name)


def build_pdf(blocks, out_path: str, doc_title: str):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                    Preformatted, Spacer, Table, TableStyle)

    _register_fonts()

    MARGIN = 1.05 * cm
    TOP = 0.95 * cm
    BOTTOM = 0.95 * cm
    PAGE_W, PAGE_H = A4
    AVAIL = PAGE_W - 2 * MARGIN
    DARK = colors.HexColor("#111111")

    body = ParagraphStyle("body", fontName="DJV", fontSize=8.1, leading=9.8,
                          spaceBefore=0, spaceAfter=1.3, alignment=4, textColor=DARK)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=8, spaceAfter=1.0)
    bullet2 = ParagraphStyle("bullet2", parent=bullet, leftIndent=17, spaceAfter=0.8)
    quote = ParagraphStyle("quote", parent=body, leftIndent=7, rightIndent=4,
                           fontSize=7.7, leading=9.1, textColor=colors.HexColor("#333333"),
                           spaceBefore=0.8, spaceAfter=1.4)
    quote_b = ParagraphStyle("quote_b", parent=quote, leftIndent=15)
    code = ParagraphStyle("code", fontName="DJV-M", fontSize=6.4, leading=7.3,
                          textColor=colors.HexColor("#101010"))
    h1 = ParagraphStyle("h1", fontName="DJV-B", fontSize=11.2, leading=13.0, spaceBefore=6.6,
                        spaceAfter=2.4, textColor=colors.HexColor("#0B3D66"), keepWithNext=1)
    h2 = ParagraphStyle("h2", fontName="DJV-B", fontSize=9.6, leading=11.2, spaceBefore=5.2,
                        spaceAfter=1.9, textColor=colors.HexColor("#12456F"), keepWithNext=1)
    h3 = ParagraphStyle("h3", fontName="DJV-B", fontSize=8.9, leading=10.4, spaceBefore=4.0,
                        spaceAfter=1.5, textColor=colors.HexColor("#1B4F72"), keepWithNext=1)
    title_s = ParagraphStyle("title", fontName="DJV-B", fontSize=13.0, leading=14.8,
                             spaceAfter=3.0, textColor=colors.HexColor("#062B4A"))

    def markup(text, base_size, mono_delta=0.3):
        pieces = []
        for t, b, c, it in inline_parts(text):
            t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if c:
                t = '<font face="DJV-M" size="%.1f">%s</font>' % (base_size - mono_delta, t)
            if b and not it:
                t = "<b>%s</b>" % t
            elif b and it:
                t = "<b><i>%s</i></b>" % t
            elif it:
                t = "<i>%s</i>" % t
            pieces.append(t)
        return "".join(pieces)

    def para(text, style):
        return Paragraph(markup(text, style.fontSize), style)

    def code_block(lines):
        tbl = Table([[Preformatted("\n".join(lines) if lines else " ", code)]], colWidths=[AVAIL])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F6F8")),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DBE0")),
            ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2.0),
            ("TOPPADDING", (0, 0), (-1, -1), 2.2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
        ]))
        return tbl

    def table_block(rows):
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        fs, lead = {1: (7.2, 8.4), 2: (7.2, 8.4), 3: (7.0, 8.2), 4: (6.6, 7.8),
                    5: (6.2, 7.4)}.get(ncols, (5.8, 7.0))
        widths = None
        while widths is None and fs > 5.0:
            widths = _col_widths(rows, ncols, fs, AVAIL)
            if widths is None:
                fs -= 0.3
                lead = fs + 1.2
        if widths is None:
            widths = [AVAIL / ncols] * ncols
        cell = ParagraphStyle("cell", fontName="DJV", fontSize=fs, leading=lead,
                              spaceBefore=0, spaceAfter=0, textColor=DARK)
        cell_h = ParagraphStyle("cellh", parent=cell, fontName="DJV-B")
        data = [[para(c, cell_h if ri == 0 else cell) for c in row]
                for ri, row in enumerate(rows)]
        tbl = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C3CAD1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7ECF1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
            ("TOPPADDING", (0, 0), (-1, -1), 1.3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.3),
        ]))
        return tbl

    story = []
    for idx, (kind, payload) in enumerate(blocks):
        if kind == "h":
            lvl, text = payload
            if idx == 0:
                story.append(para(text, title_s))
            else:
                story.append(para(text, {1: h1, 2: h2}.get(lvl, h3)))
        elif kind == "p":
            story.append(para(payload, body))
        elif kind == "quote":
            for line in payload:
                if line.startswith("- "):
                    story.append(Paragraph("•&nbsp;&nbsp;" + markup(line[2:], quote_b.fontSize),
                                           quote_b))
                elif line:
                    story.append(para(line, quote))
        elif kind == "ul":
            for indent, item in payload:
                st = bullet2 if indent >= 2 else bullet
                story.append(Paragraph("•&nbsp;&nbsp;" + markup(item, st.fontSize), st))
        elif kind == "ol":
            for num, (indent, item) in enumerate(payload, start=1):
                st = bullet2 if indent >= 2 else bullet
                story.append(Paragraph("<b>%d.</b>&nbsp;&nbsp;%s" % (num, markup(item, st.fontSize)),
                                       st))
        elif kind == "code":
            story.append(Spacer(1, 1.0))
            story.append(code_block(payload))
            story.append(Spacer(1, 1.6))
        elif kind == "table":
            story.append(Spacer(1, 0.8))
            story.append(table_block(payload))
            story.append(Spacer(1, 1.8))

    pages = {"n": 0}

    def on_page(canv, doc):
        pages["n"] += 1
        canv.saveState()
        canv.setFont("DJV", 6.2)
        canv.setFillColor(colors.HexColor("#8A949E"))
        canv.drawString(MARGIN, 0.48 * cm, doc_title[:120])
        canv.drawRightString(PAGE_W - MARGIN, 0.48 * cm, "%d" % doc.page)
        canv.setStrokeColor(colors.HexColor("#DDE2E7"))
        canv.setLineWidth(0.4)
        canv.line(MARGIN, 0.70 * cm, PAGE_W - MARGIN, 0.70 * cm)
        canv.restoreState()

    doc = BaseDocTemplate(out_path, pagesize=A4,
                          leftMargin=MARGIN, rightMargin=MARGIN,
                          topMargin=TOP, bottomMargin=BOTTOM,
                          title=doc_title, author="MPMC Module 2 notes")
    frame = Frame(MARGIN, BOTTOM, AVAIL, PAGE_H - TOP - BOTTOM, id="body",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame], onPage=on_page)])
    doc.build(story)
    return pages["n"]


# --------------------------------------------------------------------------- #
# DOCX (python-docx)
# --------------------------------------------------------------------------- #

def build_docx(blocks, out_path: str, doc_title: str):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    BODY_FONT, MONO_FONT = "Calibri", "Consolas"
    BODY_SIZE = 8.5
    CONTENT_W = 18.9  # cm

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.05)
    sec.top_margin = Cm(0.9)
    sec.bottom_margin = Cm(0.9)
    sec.footer_distance = Cm(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_after = Pt(1.4)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    for name, size, color, sb, sa in (("Title", 14.0, "#062B4A", 0, 4),
                                      ("Heading 1", 11.8, "#0B3D66", 7, 2.5),
                                      ("Heading 2", 10.0, "#12456F", 5.5, 2),
                                      ("Heading 3", 9.1, "#1B4F72", 4.5, 1.5)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        st.paragraph_format.space_before = Pt(sb)
        st.paragraph_format.space_after = Pt(sa)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.0

    def shade(tc_or_pPr, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_or_pPr.append(shd)

    def add_runs(p, text, size=BODY_SIZE, base_bold=False, italic=False):
        for t, b, c, it in inline_parts(text):
            r = p.add_run(t)
            r.bold = b or base_bold
            r.italic = it or italic
            r.font.size = Pt(size - 0.3 if c else size)
            r.font.name = MONO_FONT if c else BODY_FONT
        return p

    def add_para(text, size=BODY_SIZE, indent=0.0, style=None, italic=False,
                 space_after=1.4, prefix=None, prefix_bold=False):
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.0
        if indent:
            pf.left_indent = Cm(indent)
        if prefix:
            r = p.add_run(prefix)
            r.bold = prefix_bold
            r.font.size = Pt(size)
            r.font.name = BODY_FONT
            r.italic = italic
        add_runs(p, text, size=size, italic=italic)
        return p

    def fixed_layout(tbl):
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl._tbl.tblPr.append(layout)

    def cell_margins(tbl, twips=22):
        mar = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement("w:" + side)
            el.set(qn("w:w"), str(twips))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tbl._tbl.tblPr.append(mar)

    def add_code(lines):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Cm(CONTENT_W)
        shade(cell._tc.get_or_add_tcPr(), "F4F6F8")
        first = True
        for ln in (lines or [""]):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            r = p.add_run(ln if ln else " ")
            r.font.name = MONO_FONT
            r.font.size = Pt(6.8)
        fixed_layout(tbl)
        cell_margins(tbl, 40)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0.5)
        spacer.paragraph_format.space_before = Pt(0)
        for r in spacer.runs:
            r.font.size = Pt(2)

    def add_table(rows):
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        size = {1: 7.4, 2: 7.4, 3: 7.2, 4: 6.8, 5: 6.4}.get(ncols, 6.0)
        pt_widths = None
        while pt_widths is None and size > 5.0:
            pt_widths = _col_widths(rows, ncols, size, CONTENT_W * 28.3465)
            if pt_widths is None:
                size -= 0.3
        if pt_widths is None:
            pt_widths = [CONTENT_W * 28.3465 / ncols] * ncols
        tbl = doc.add_table(rows=len(rows), cols=ncols)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl.autofit = False
        for ri, row in enumerate(rows):
            for ci, txt in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.width = Cm(pt_widths[ci] * 2.54 / 72.0)
                p = cell.paragraphs[0]
                pf = p.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.0
                add_runs(p, txt, size=size, base_bold=(ri == 0))
                if ri == 0:
                    shade(cell._tc.get_or_add_tcPr(), "E7ECF1")
        fixed_layout(tbl)
        cell_margins(tbl, 22)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0.5)
        spacer.paragraph_format.space_before = Pt(0)

    for idx, (kind, payload) in enumerate(blocks):
        if kind == "h":
            lvl, text = payload
            if idx == 0:
                p = doc.add_paragraph(style="Title")
                add_runs(p, text, size=14.0, base_bold=True)
            else:
                style = {1: "Heading 1", 2: "Heading 2"}.get(lvl, "Heading 3")
                p = doc.add_paragraph(style=style)
                add_runs(p, text, size={1: 11.8, 2: 10.0}.get(lvl, 9.1), base_bold=True)
        elif kind == "p":
            add_para(payload)
        elif kind == "quote":
            for line in payload:
                if line.startswith("- "):
                    add_para(line[2:], size=7.9, indent=0.75, italic=True,
                             space_after=1.0, prefix="•  ")
                elif line:
                    add_para(line, size=7.9, indent=0.32, italic=True, space_after=1.6)
        elif kind == "ul":
            for indent, item in payload:
                add_para(item, indent=0.42 if indent < 2 else 0.85,
                         space_after=1.0, prefix="•  ")
        elif kind == "ol":
            for k, (indent, item) in enumerate(payload, start=1):
                add_para(item, indent=0.42 if indent < 2 else 0.85,
                         space_after=1.0, prefix="%d.  " % k, prefix_bold=True)
        elif kind == "code":
            add_code(payload)
        elif kind == "table":
            add_table(payload)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run(doc_title[:110] + "   |   page ")
    run.font.size = Pt(6.5)
    run.font.name = BODY_FONT
    run.font.color.rgb = RGBColor.from_string("7A828A")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "13")
    rpr.append(sz)
    r.append(rpr)
    fld.append(r)
    fp._p.append(fld)

    doc.save(out_path)


# --------------------------------------------------------------------------- #

def convert(md_path: str):
    with open(md_path, encoding="utf-8") as fh:
        text = fh.read()
    blocks = parse_blocks(text)
    base = os.path.splitext(md_path)[0]
    title = next((plain(p[1]) for k, p in blocks if k == "h"), os.path.basename(base))
    npages = build_pdf(blocks, base + ".pdf", title)
    build_docx(blocks, base + ".docx", title)
    print("%-52s -> %2d pages (pdf) + docx" % (os.path.basename(md_path), npages))
    return npages


if __name__ == "__main__":
    for path in sys.argv[1:]:
        convert(path)
